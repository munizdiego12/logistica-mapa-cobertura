"""
Fonte de pedidos (Etapa 7): interface única para trazer um lote de pedidos de qualquer origem, o modelo padrão de CSV/XLSX e a
validação linha a linha.

Camadas:
    0) `data_entrega` é OPCIONAL e só informativa (decisão do operador: pedidos sem filtro de data; o id identifica o pedido e o
       operador consulta o dia na base dele): se vier preenchida o formato é validado e ela é preservada no lote, mas nada filtra por ela;
    1) uma FONTE lê a origem e devolve um `LoteBruto` (cada linha com os campos padrão, ainda como texto): hoje só `FonteArquivo`
       (CSV/XLSX enviado pelo operador). Google Sheets, Superset e BigQuery (Etapa 15) entram como novas classes com o método
       `ler()`, registradas em `FONTES`, e reaproveitam a validação;
    2) `validar_lote` valida e padroniza cada linha e devolve um `RelatorioLote` (válidas, erros, avisos) para o operador conferir
       ANTES de aceitar o lote. Nada é gravado: o lote é temporário, vive só na resposta e na tela, e é descartado depois do uso.

Dado pessoal: cliente e endereço NUNCA vão para log nem para mensagem de erro. Os motivos dos erros citam só o NOME do campo e o
número da linha, jamais o valor da célula; erros de leitura do arquivo usam textos fixos.
"""
import csv
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from typing import Optional, Protocol

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import casamento_lojas
import lojas as modulo_lojas

MAX_BYTES = 10 * 1024 * 1024
MAX_LINHAS = 20000
MAX_ERROS_NA_RESPOSTA = 1000
MAX_LINHAS_POR_AVISO = 20

# Colunas do modelo padrão, na ordem do arquivo.
CAMPOS = (
    "id", "loja", "cliente", "endereco", "numero", "complemento", "bairro", "cidade", "uf", "cep",
    "data_entrega", "janela_inicio", "janela_fim", "volume", "observacao",
)
OBRIGATORIOS = ("id", "endereco", "numero", "cidade", "uf")
LIMITES = {"id": 60, "loja": 120, "cliente": 150, "endereco": 250, "numero": 30, "complemento": 120, "bairro": 120, "cidade": 120, "observacao": 500}

# Cabeçalhos reconhecidos POR NOME, sem diferença de maiúsculas, acentos, espaços nem separadores ("ID do pedido" = "id_do_pedido").
_ALIASES = {
    "id": ("id", "idpedido", "iddopedido", "pedido", "numeropedido", "numerodopedido", "npedido", "codigopedido", "codigodopedido", "order", "orderid"),
    "loja": ("loja", "nomedaloja", "nomeloja", "lojaorigem", "lojadeorigem", "codigoloja", "codigodaloja", "codigoexterno", "nomeoperacao",
             "nomedeoperacao", "filial", "hub", "store"),
    "cliente": ("cliente", "nomecliente", "nomedocliente", "destinatario", "customer"),
    "endereco": ("endereco", "rua", "logradouro", "address", "street", "enderecodeentrega"),
    "numero": ("numero", "num", "nro", "number"),
    "complemento": ("complemento", "compl"),
    "bairro": ("bairro", "neighborhood", "district"),
    "cidade": ("cidade", "city", "municipio"),
    "uf": ("uf", "estado", "state"),
    "cep": ("cep", "zipcode", "postalcode", "codigopostal"),
    "data_entrega": ("dataentrega", "datadeentrega", "dataprevista", "dataprevistadeentrega", "data", "entrega"),
    "janela_inicio": ("janelainicio", "janelainicial", "iniciodajanela", "horainicio", "horarioinicio", "inicio", "janeladeentregainicio"),
    "janela_fim": ("janelafim", "janelafinal", "fimdajanela", "horafim", "horariofim", "fim", "janeladeentregafim"),
    "volume": ("volume", "vol", "quantidade", "pedidos", "qtd"),
    "observacao": ("observacao", "observacoes", "obs", "nota", "notas"),
}


class ErroDeArquivo(ValueError):
    """Problema com o arquivo todo (formato, tamanho, colunas). A mensagem é fixa, em linguagem do operador, sem conteúdo do arquivo."""


# ----------------------------------------------------------------------------------------------------------
# Interface única
# ----------------------------------------------------------------------------------------------------------
@dataclass
class LinhaBruta:
    numero: int  # linha na origem (1 = cabeçalho)
    campos: dict  # campo padrão -> texto (campos que a origem não tem ficam ausentes)


@dataclass
class LoteBruto:
    origem: str  # "arquivo" (depois: "sheets", "superset", "bigquery")
    formato: str  # "csv", "xlsx"...
    colunas: dict  # campo padrão -> nome do cabeçalho na origem
    linhas: list = field(default_factory=list)

    @property
    def tem_lojas(self) -> bool:
        """A origem tem a coluna loja e pelo menos uma linha a preencheu (só então é preciso conferir o cadastro)."""
        return "loja" in self.colunas and any(l.campos.get("loja") for l in self.linhas)


class FontePedidos(Protocol):
    """Contrato de toda fonte de pedidos: `ler()` devolve o lote já com os nomes de campo padrão."""

    def ler(self) -> LoteBruto: ...


def _normalizar_cabecalho(texto) -> str:
    return modulo_lojas.normalizar_cabecalho(texto).replace("_", "")


def _reconhecer(cabecalho: list) -> dict:
    """campo padrão -> posição da coluna, pelo nome do cabeçalho (o primeiro que casar)."""
    posicoes = {}
    for i, c in enumerate(cabecalho):
        posicoes.setdefault(_normalizar_cabecalho(c), i)
    achados = {}
    for campo, aliases in _ALIASES.items():
        for alias in aliases:
            if alias in posicoes:
                achados[campo] = posicoes[alias]
                break
    return achados


def _texto(valor) -> str:
    """Célula como texto limpo: números inteiros sem '.0', datas em ISO, horas como HH:MM."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else str(valor)
    if isinstance(valor, datetime):
        return valor.date().isoformat() if valor.time() == time(0, 0) else valor.isoformat(sep=" ")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    return " ".join(str(valor).split())


def _montar_lote(formato: str, cabecalho: list, linhas_brutas: list) -> LoteBruto:
    posicoes = _reconhecer(cabecalho)
    lote = LoteBruto(origem="arquivo", formato=formato, colunas={c: str(cabecalho[i]).strip() for c, i in posicoes.items()})
    for numero, celulas in linhas_brutas:
        textos = [_texto(c) for c in celulas]
        if not any(textos):
            continue
        campos = {c: (textos[i] if i < len(textos) else "") for c, i in posicoes.items()}
        lote.linhas.append(LinhaBruta(numero, campos))
        if len(lote.linhas) > MAX_LINHAS:
            raise ErroDeArquivo(f"O arquivo tem mais de {MAX_LINHAS:,} linhas".replace(",", ".") + ". Divida em lotes menores.")
    return lote


def _ler_csv(conteudo: bytes) -> LoteBruto:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252", errors="replace")
    melhor = None
    for delimitador in (";", ",", "\t"):
        leitor = csv.reader(io.StringIO(texto, newline=""), delimiter=delimitador)
        try:
            linhas = [(leitor.line_num, linha) for linha in leitor]
        except csv.Error:  # célula gigante, aspas quebradas...: texto fixo, nunca o conteúdo
            raise ErroDeArquivo("Não consegui ler o CSV. Confira o arquivo ou use o modelo padrão (botões CSV e XLSX).") from None
        cabecalho_idx = next((i for i, (_, l) in enumerate(linhas) if any(c.strip() for c in l)), None)
        if cabecalho_idx is None:
            continue
        pontos = len(_reconhecer(linhas[cabecalho_idx][1]))
        if melhor is None or pontos > melhor[0]:
            melhor = (pontos, linhas, cabecalho_idx)
    if melhor is None:
        raise ErroDeArquivo("O arquivo está vazio. Use o modelo padrão (botões CSV e XLSX).")
    _, linhas, cabecalho_idx = melhor
    return _montar_lote("csv", linhas[cabecalho_idx][1], linhas[cabecalho_idx + 1:])


def _ler_xlsx(conteudo: bytes) -> LoteBruto:
    try:
        pasta = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:  # arquivo corrompido ou não é um .xlsx: texto fixo, nunca o conteúdo
        raise ErroDeArquivo("Não consegui ler a planilha. Confira se é um arquivo .xlsx válido ou use o modelo padrão.") from None
    melhor = None
    try:
        for aba in pasta.worksheets:
            linhas = [(i, list(l)) for i, l in enumerate(aba.iter_rows(values_only=True), start=1)]
            for posicao, (_, celulas) in enumerate(linhas[:10]):  # o cabeçalho pode ter títulos acima
                pontos = len(_reconhecer([_texto(c) for c in celulas]))
                if melhor is None or pontos > melhor[0]:
                    melhor = (pontos, linhas, posicao)
    except Exception:
        raise ErroDeArquivo("Não consegui ler a planilha. Confira se é um arquivo .xlsx válido ou use o modelo padrão.") from None
    finally:
        pasta.close()
    if melhor is None:
        raise ErroDeArquivo("A planilha está vazia. Use o modelo padrão (botões CSV e XLSX).")
    _, linhas, posicao = melhor
    return _montar_lote("xlsx", [_texto(c) for c in linhas[posicao][1]], linhas[posicao + 1:])


class FonteArquivo:
    """CSV ou XLSX enviado pelo operador (o modelo padrão, o modelo antigo ou variações dos cabeçalhos)."""

    def __init__(self, conteudo: bytes, nome: str = ""):
        self.conteudo = conteudo
        self.nome = (nome or "").lower()

    def ler(self) -> LoteBruto:
        if len(self.conteudo) > MAX_BYTES:
            raise ErroDeArquivo("O arquivo é grande demais (máximo de 10 MB). Divida em lotes menores.")
        if self.nome.endswith(".csv"):
            return _ler_csv(self.conteudo)
        if self.nome.endswith(".xlsx"):
            return _ler_xlsx(self.conteudo)
        if self.nome.endswith(".xls"):
            raise ErroDeArquivo("Planilhas .xls antigas não são aceitas: salve como .xlsx ou .csv.")
        raise ErroDeArquivo("Formato de arquivo inválido. Envie um arquivo .csv ou .xlsx.")


FONTES = {"arquivo": FonteArquivo}  # Etapa 15: "sheets", "superset", "bigquery"


def carregar_lote(fonte: FontePedidos, lojas_cadastradas: Optional[list] = None) -> "RelatorioLote":
    """Interface única: lê a fonte e valida o lote. `lojas_cadastradas` só é necessário se o lote tiver a coluna loja preenchida."""
    return validar_lote(fonte.ler(), lojas_cadastradas)


# ----------------------------------------------------------------------------------------------------------
# Valores
# ----------------------------------------------------------------------------------------------------------
def interpretar_cep(texto: str) -> tuple:
    """(CEP no formato 01310-100, aviso ou None). ValueError se não tiver 8 dígitos (7, sem o zero inicial do Excel, é completado)."""
    digitos = re.sub(r"\D", "", texto)
    if re.search(r"[A-Za-z]", texto) or not digitos:
        raise ValueError("cep")
    aviso = None
    if len(digitos) == 7:
        digitos, aviso = "0" + digitos, "CEP com 7 dígitos completado com um zero à esquerda"
    if len(digitos) != 8 or digitos == "00000000":
        raise ValueError("cep")
    return f"{digitos[:5]}-{digitos[5:]}", aviso


def interpretar_volume(texto: str) -> int:
    try:
        numero = modulo_lojas._numero_ptbr(texto)
    except ValueError:
        raise ValueError("volume") from None
    if numero != int(numero) or not 1 <= numero <= 9999:
        raise ValueError("volume")
    return int(numero)


_RE_DATA_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?$")
_RE_DATA_BR = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})$")


def interpretar_data(texto: str) -> str:
    """'2026-10-20', '20/10/2026', '20-10-26', '2026-10-20 00:00:00' -> '2026-10-20'. ValueError se não for uma data de verdade."""
    achou = _RE_DATA_ISO.match(texto)
    if achou:
        ano, mes, dia = (int(g) for g in achou.groups())
    else:
        achou = _RE_DATA_BR.match(texto)
        if not achou:
            raise ValueError("data")
        dia, mes, ano = int(achou.group(1)), int(achou.group(2)), int(achou.group(3))
        ano = ano + 2000 if ano < 100 else ano
    try:
        data = date(ano, mes, dia)
    except ValueError:
        raise ValueError("data") from None
    if not 2000 <= ano <= 2100:
        raise ValueError("data")
    return data.isoformat()


_RE_HORA = re.compile(r"^(?:\d{4}-\d{2}-\d{2}[ T])?(\d{1,2})(?::(\d{2})(?::\d{2})?|h(\d{2})?)?$")


def interpretar_hora(texto: str) -> str:
    """'8', '8:30', '08:00:00', '8h', '8h30' -> 'HH:MM'. ValueError se não for um horário."""
    achou = _RE_HORA.match(texto.lower().replace(" h", "h"))
    if not achou:
        raise ValueError("hora")
    hora, minuto = int(achou.group(1)), int(achou.group(2) or achou.group(3) or 0)
    if hora > 23 or minuto > 59:
        raise ValueError("hora")
    return f"{hora:02d}:{minuto:02d}"


# ----------------------------------------------------------------------------------------------------------
# Validação linha a linha
# ----------------------------------------------------------------------------------------------------------
@dataclass
class RelatorioLote:
    origem: str
    formato: str
    colunas: dict
    colunas_opcionais_ausentes: list
    total_linhas: int
    pedidos: list  # só as válidas
    erros: list  # [{linha, id, motivos}]
    avisos: list  # [{motivo, quantidade, linhas}]
    erros_por_motivo: dict
    lojas: list  # [{loja_id, nome, quantidade}] das válidas
    sem_loja: int
    com_aviso: int
    erros_total: int

    def para_dict(self) -> dict:
        return {
            "origem": self.origem,
            "formato": self.formato,
            "colunas": self.colunas,
            "colunas_opcionais_ausentes": self.colunas_opcionais_ausentes,
            "resumo": {
                "linhas": self.total_linhas, "validas": len(self.pedidos), "com_erro": self.erros_total, "com_aviso": self.com_aviso,
                "sem_loja": self.sem_loja,
            },
            "pedidos": self.pedidos,
            "erros": self.erros,
            "erros_truncados": self.erros_total > len(self.erros),
            "erros_por_motivo": self.erros_por_motivo,
            "avisos": self.avisos,
            "lojas": self.lojas,
        }


MOTIVO_SEM_CEP = "sem CEP: o endereço será localizado só pelo texto e pode ficar aproximado"
MOTIVO_SEM_LOJA = "linha sem loja: vale para a loja escolhida na tela"
MOTIVO_LOJAS_NAO_CONFERIDAS = "as lojas do arquivo não foram conferidas com o cadastro"


def validar_lote(bruto: LoteBruto, lojas_cadastradas: Optional[list] = None) -> RelatorioLote:
    """Valida e padroniza cada linha. Os motivos citam só o nome do campo (nunca o valor da célula)."""
    faltam = [c for c in OBRIGATORIOS if c not in bruto.colunas]
    if faltam:
        raise ErroDeArquivo(
            "Faltam colunas obrigatórias no cabeçalho: " + ", ".join(faltam)
            + ". Use o modelo padrão (botões CSV e XLSX) ou acrescente as colunas."
        )
    conferir_lojas = lojas_cadastradas is not None
    pedidos, erros, avisos_por_motivo = [], [], {}
    primeira_linha_do_id = {}
    erros_por_motivo = Counter()
    lojas_contagem = Counter()
    sem_loja = com_aviso = 0

    def avisar(motivo, linha):
        entrada = avisos_por_motivo.setdefault(motivo, {"motivo": motivo, "quantidade": 0, "linhas": []})
        entrada["quantidade"] += 1
        if len(entrada["linhas"]) < MAX_LINHAS_POR_AVISO:
            entrada["linhas"].append(linha)

    for linha in bruto.linhas:
        c = linha.campos
        valor = lambda campo: c.get(campo, "")  # noqa: E731
        motivos, avisos = [], []

        vazios = [campo for campo in OBRIGATORIOS if not valor(campo)]
        if vazios:
            motivos.append("campo obrigatório vazio: " + ", ".join(vazios))
        longos = [campo for campo, limite in LIMITES.items() if len(valor(campo)) > limite]
        if longos:
            motivos.append("texto longo demais em: " + ", ".join(longos))

        id_pedido = valor("id")
        if id_pedido:
            chave_id = id_pedido.lower()
            anterior = primeira_linha_do_id.setdefault(chave_id, linha.numero)
            if anterior != linha.numero:
                motivos.append(f"id repetido (a primeira vez está na linha {anterior})")

        uf = None
        if valor("uf"):
            try:
                uf = modulo_lojas.normalizar_uf(valor("uf"))
            except ValueError:
                motivos.append("UF inválida (use a sigla, ex.: SP)")

        cep = None
        if valor("cep"):
            try:
                cep, aviso = interpretar_cep(valor("cep"))
                if aviso:
                    avisos.append(aviso)
            except ValueError:
                motivos.append("CEP inválido (precisa ter 8 dígitos)")
        else:
            avisos.append(MOTIVO_SEM_CEP)

        volume = 1
        if valor("volume"):
            try:
                volume = interpretar_volume(valor("volume"))
            except ValueError:
                motivos.append("volume inválido (use um número inteiro de 1 a 9999)")

        data_entrega = None  # só informativa: validada e preservada, nunca filtra nem muda a roteirização
        if valor("data_entrega"):
            try:
                data_entrega = interpretar_data(valor("data_entrega"))
            except ValueError:
                motivos.append("data de entrega inválida (use dd/mm/aaaa ou aaaa-mm-dd)")

        inicio = fim = None
        horario_ok = True
        for campo, rotulo in (("janela_inicio", "início"), ("janela_fim", "fim")):
            if valor(campo):
                try:
                    resultado = interpretar_hora(valor(campo))
                    inicio, fim = (resultado, fim) if campo == "janela_inicio" else (inicio, resultado)
                except ValueError:
                    horario_ok = False
                    motivos.append(f"horário de {rotulo} da janela inválido (use hh:mm)")
        if horario_ok:
            if bool(valor("janela_inicio")) != bool(valor("janela_fim")):
                motivos.append("janela incompleta: informe o início e o fim")
            elif inicio and fim and inicio >= fim:
                motivos.append("janela invertida: o início precisa ser antes do fim")

        loja_id = loja_nome = None
        if "loja" in bruto.colunas:
            if not valor("loja"):
                avisos.append(MOTIVO_SEM_LOJA)
            elif conferir_lojas:
                situacao, loja = casamento_lojas.casar_loja(valor("loja"), lojas_cadastradas)
                if situacao == "ok":
                    loja_id, loja_nome = int(loja["id"]), loja["nome"]
                else:
                    motivos.append({
                        "nao_encontrada": "loja não encontrada no cadastro",
                        "ambigua": "loja ambígua: o nome casa com mais de uma loja",
                        "desativada": "loja desativada no cadastro",
                    }[situacao])
            else:
                avisos.append(MOTIVO_LOJAS_NAO_CONFERIDAS)

        if motivos:
            erros_por_motivo.update(m.split(":")[0].split(" (")[0] for m in motivos)
            erros.append({"linha": linha.numero, "id": id_pedido[:60], "motivos": motivos})
            continue

        for aviso in avisos:
            avisar(aviso, linha.numero)
        com_aviso += 1 if avisos else 0
        if loja_id is not None:
            lojas_contagem[(loja_id, loja_nome)] += 1
        elif "loja" in bruto.colunas:
            sem_loja += 1
        pedidos.append({
            "linha": linha.numero, "id": id_pedido, "loja": valor("loja") or None, "loja_id": loja_id, "loja_nome": loja_nome,
            "cliente": valor("cliente") or None, "endereco": valor("endereco"), "numero": valor("numero"),
            "complemento": valor("complemento") or None, "bairro": valor("bairro") or None, "cidade": valor("cidade"), "uf": uf,
            "cep": cep, "data_entrega": data_entrega, "janela_inicio": inicio, "janela_fim": fim, "volume": volume,
            "observacao": valor("observacao") or None, "avisos": avisos,
        })

    return RelatorioLote(
        origem=bruto.origem, formato=bruto.formato, colunas=bruto.colunas,
        colunas_opcionais_ausentes=[c for c in CAMPOS if c not in bruto.colunas and c not in OBRIGATORIOS],
        total_linhas=len(bruto.linhas), pedidos=pedidos, erros=erros[:MAX_ERROS_NA_RESPOSTA], avisos=list(avisos_por_motivo.values()),
        erros_por_motivo=dict(erros_por_motivo), erros_total=len(erros),
        lojas=[{"loja_id": i, "nome": n, "quantidade": q} for (i, n), q in sorted(lojas_contagem.items(), key=lambda x: x[0][1].lower())],
        sem_loja=sem_loja, com_aviso=com_aviso,
    )


# ----------------------------------------------------------------------------------------------------------
# Modelo padrão (CSV e XLSX)
# ----------------------------------------------------------------------------------------------------------
def _exemplos() -> list:
    amanha = (datetime.now(timezone.utc) - timedelta(hours=3) + timedelta(days=1)).date().isoformat()
    base = [
        ["1001", "", "Cliente Exemplo 1", "Avenida Paulista", "1500", "Sala 10", "Bela Vista", "Sao Paulo", "SP", "01310-100", amanha, "08:00", "13:00", 1, ""],
        ["1002", "", "Cliente Exemplo 2", "Rua Augusta", "1500", "", "Consolacao", "Sao Paulo", "SP", "01304-001", amanha, "13:00", "18:00", 1, "Deixar com o porteiro"],
        ["1003", "", "Cliente Exemplo 3", "Rua Oscar Freire", "850", "", "Cerqueira Cesar", "Sao Paulo", "SP", "01426-000", amanha, "", "", 2, ""],
        ["1004", "", "Cliente Exemplo 4", "Rua dos Pinheiros", "600", "Apto 42", "Pinheiros", "Sao Paulo", "SP", "05422-001", amanha, "18:00", "21:00", 1, ""],
        ["1005", "", "Cliente Exemplo 5", "Alameda dos Maracatins", "650", "", "Moema", "Sao Paulo", "SP", "04089-001", amanha, "", "", "", ""],
    ]
    return base


INSTRUCOES = [
    "Modelo padrão de pedidos",
    "",
    "Preencha a aba Modelo_Pedidos, uma linha por pedido. Apague as linhas de exemplo antes de enviar.",
    "Obrigatórias: id, endereco, numero, cidade e uf. As demais colunas são opcionais.",
    "id: código do pedido; não pode se repetir no arquivo.",
    "loja: nome de operação, código externo ou nome da loja cadastrada. Se a coluna ficar em branco, o lote vale para a loja escolhida na tela.",
    "numero: use S/N quando o endereço não tiver número.",
    "cep: 8 dígitos, com ou sem hífen (ex.: 01310-100). Ajuda a localizar o endereço.",
    "uf: sigla do estado (ex.: SP).",
    "data_entrega: dd/mm/aaaa ou aaaa-mm-dd.",
    "janela_inicio e janela_fim: horários hh:mm (ex.: 08:00 e 13:00); informe os dois ou nenhum, e o início precisa ser antes do fim.",
    "volume: número inteiro de 1 a 9999; vale 1 se ficar em branco.",
    "Antes de aceitar o lote, a tela mostra as linhas com erro e o motivo; você pode aceitar só as linhas válidas.",
    "O arquivo não é guardado: o lote é descartado depois do uso.",
]


def gerar_modelo_csv() -> bytes:
    """Modelo padrão em CSV (UTF-8 com BOM e ';', que o Excel em português abre direito)."""
    saida = io.StringIO(newline="")
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(CAMPOS)
    escritor.writerows(_exemplos())
    return b"\xef\xbb\xbf" + saida.getvalue().encode("utf-8")


def gerar_modelo_xlsx() -> bytes:
    """Modelo padrão em XLSX: aba Modelo_Pedidos (cabeçalho e 5 exemplos) e aba Instruções."""
    pasta = openpyxl.Workbook()
    aba = pasta.active
    aba.title = "Modelo_Pedidos"
    aba.append(list(CAMPOS))
    for linha in _exemplos():
        aba.append(linha)
    preenchimento = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fonte = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    obrigatoria = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    borda = Border(*(Side(style="thin", color="CBD5E1") for _ in range(4)))
    for coluna, nome in enumerate(CAMPOS, start=1):
        celula = aba.cell(row=1, column=coluna)
        celula.fill = obrigatoria if nome in OBRIGATORIOS else preenchimento
        celula.font = fonte
        celula.alignment = Alignment(horizontal="center", vertical="center")
    for linha in aba.iter_rows(min_row=2, max_row=aba.max_row):
        for celula in linha:
            celula.border = borda
            celula.number_format = "@"  # texto: preserva o zero do CEP e o formato do horário
    for coluna in aba.columns:
        tamanho = max(len(str(c.value or "")) for c in coluna)
        aba.column_dimensions[openpyxl.utils.get_column_letter(coluna[0].column)].width = max(tamanho + 4, 12)
    aba.freeze_panes = "A2"
    guia = pasta.create_sheet("Instruções")
    for texto in INSTRUCOES:
        guia.append([texto])
    guia["A1"].font = Font(name="Calibri", size=14, bold=True)
    guia.column_dimensions["A"].width = 120
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()
