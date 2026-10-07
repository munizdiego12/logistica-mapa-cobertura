"""Testes do endpoint de cobertura e da exportação XLSX, com o banco e o geocodificador simulados."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

import main
from cobertura import ATRIBUICAO_CNEFE, LEGENDA_COBERTURA

REQ = {"origem_rua": "SQS 108", "origem_num": "1", "origem_cep": "70347000", "raio_km": 30.0}


def _prefixo(prefixo, distancia, dispersao, uf="DF", localidade="ASA SUL", pct=100.0):
    return {
        "prefixo": prefixo, "uf": uf, "cod_municipio": 5300108, "cidade": "Brasília", "localidade": localidade,
        "localidade_pct": pct, "lat": -15.82, "lon": -47.90, "n_enderecos": 100,
        "dispersao_km": dispersao, "distancia_km": distancia,
    }


def _faixa(uf, cidade, distancia, ibge):
    return {
        "ibge": ibge, "uf": uf, "cidade": cidade, "bairro": "Sede", "cep_inicial": "72800000",
        "cep_final": "72899999", "faixa_completa": "72800-000 a 72899-999", "distancia_km": distancia,
        "dias_sla": 2, "lat": -16.25, "lon": -47.95,
    }


@pytest.fixture
def cliente(monkeypatch):
    async def geocode_falso(client, *args, **kwargs):
        return -15.8198, -47.9012, "Brasília", "DF", "70347000", "Asa Sul", "VALIDO"

    async def ibge_falso(uf, cidade):
        return 5300108

    monkeypatch.setattr(main, "geocode_async", geocode_falso)
    monkeypatch.setattr(main.database, "buscar_ibge_municipio", ibge_falso)
    return TestClient(main.app)


def _simular_banco(monkeypatch, prefixos, ufs, faixas):
    async def _prefixos(lat, lon, raio):
        return prefixos

    async def _ufs():
        return ufs

    async def _faixas(lat, lon, raio):
        return faixas

    monkeypatch.setattr(main.database, "consultar_prefixos_por_raio", _prefixos)
    monkeypatch.setattr(main.database, "ufs_com_prefixos", _ufs)
    monkeypatch.setattr(main.database, "consultar_ceps_por_raio", _faixas)


def test_hub_do_df_usa_prefixos_e_mantem_o_entorno_em_go_como_parcial(cliente, monkeypatch):
    _simular_banco(
        monkeypatch,
        prefixos=[_prefixo("70254", 3.0, 0.14), _prefixo("72499", 20.0, 90.0, localidade="GAMA", pct=30.0)],
        ufs={"DF", "AC"},
        faixas=[_faixa("DF", "Brasília", 1.0, 5300108), _faixa("GO", "Luziânia", 25.0, 5212501)],
    )
    r = cliente.post("/api/cobertura-ceps", json=REQ)
    assert r.status_code == 200
    d = r.json()
    assert d["hub"]["ibge"] == 5300108  # vem de ibge_municipios, não de valor fixo
    assert d["total_pontos"] == 3 and d["aviso"] is None
    assert [(p["uf"], p["cobertura"], p["precisao"]) for p in d["pontos_cobertos"]] == [
        ("DF", "Total", "prefixo"), ("DF", "Parcial", "prefixo"), ("GO", "Parcial", "faixa")]
    assert d["pontos_cobertos"][1]["bairro"] == "Gama e outros"  # localidade com só 30% do prefixo
    assert d["resumo_cobertura"] == {"total": 1, "parcial": 2, "usa_cnefe": True}
    assert d["legenda_cobertura"] == LEGENDA_COBERTURA and d["fonte"] == ATRIBUICAO_CNEFE


def test_hub_de_sp_sem_prefixos_continua_com_as_faixas_manuais(cliente, monkeypatch):
    _simular_banco(monkeypatch, prefixos=[], ufs={"DF", "AC"},
                   faixas=[_faixa("SP", "São Paulo", float(i), 3550308) for i in range(24)])
    d = cliente.post("/api/cobertura-ceps", json=REQ).json()
    assert d["total_pontos"] == 24 and d["aviso"] is None
    assert {p["cobertura"] for p in d["pontos_cobertos"]} == {"Parcial"}
    assert d["fonte"] is None  # sem dado do CNEFE, sem atribuição ao CNEFE


def test_hub_sem_nenhuma_cobertura_devolve_aviso_como_florianopolis(cliente, monkeypatch):
    _simular_banco(monkeypatch, prefixos=[], ufs={"DF", "AC"}, faixas=[])
    d = cliente.post("/api/cobertura-ceps", json=REQ).json()
    assert d["total_pontos"] == 0 and d["pontos_cobertos"] == []
    assert d["aviso"] == (
        "Ainda não temos CEPs cadastrados para essa região. Confira o endereço da loja "
        "ou peça ao administrador do sistema para incluir a região."
    )
    assert d["resumo_cobertura"] == {"total": 0, "parcial": 0, "usa_cnefe": False}


def test_falha_do_banco_vira_503_e_nao_o_aviso_de_sem_cobertura(cliente, monkeypatch):
    async def quebrado(*args):
        raise RuntimeError("banco de dados indisponível")

    _simular_banco(monkeypatch, prefixos=[], ufs=set(), faixas=[])
    monkeypatch.setattr(main.database, "consultar_prefixos_por_raio", quebrado)
    r = cliente.post("/api/cobertura-ceps", json=REQ)
    assert r.status_code == 503
    assert "cobertura de CEPs" in r.json()["detail"]


TERMOS_TECNICOS = (
    "ceps_reais", "cep_prefixos", "ibge_municipios", "scripts", ".py", ".csv", "tabela", "banco de dados",
    "BrasilAPI", "ViaCEP", "Nominatim", "API", "SQL", "carregar_", "DATABASE",
)


def test_mensagens_de_tela_da_cobertura_nao_citam_termos_tecnicos(cliente, monkeypatch):
    mensagens = []

    # 1) aviso de região sem cobertura
    _simular_banco(monkeypatch, prefixos=[], ufs={"DF"}, faixas=[])
    mensagens.append(cliente.post("/api/cobertura-ceps", json=REQ).json()["aviso"])

    # 2) falha do banco (503)
    async def quebrado(*args):
        raise RuntimeError("banco de dados indisponível")

    monkeypatch.setattr(main.database, "consultar_prefixos_por_raio", quebrado)
    mensagens.append(cliente.post("/api/cobertura-ceps", json=REQ).json()["detail"])

    # 3) endereço da loja não encontrado (422)
    async def geocode_vazio(client, *args, **kwargs):
        return None, None, "", "", "", "", "ERRO_CEP_INVALIDO"

    monkeypatch.setattr(main, "geocode_async", geocode_vazio)
    r = cliente.post("/api/cobertura-ceps", json=REQ)
    assert r.status_code == 422
    mensagens.append(r.json()["detail"])

    assert all(mensagens) and len(mensagens) == 3
    for texto in mensagens:
        for termo in TERMOS_TECNICOS:
            assert termo.lower() not in texto.lower(), f"termo técnico '{termo}' em: {texto}"


def _ponto(cobertura, precisao="prefixo", ibge=5300108, bairro="Asa Sul"):
    return {
        "ibge": ibge, "uf": "DF", "cidade": "Brasília", "bairro": bairro, "cep_inicial": "70254000",
        "cep_final": "70254999", "distancia_km": 3.0, "dias_sla": 1, "cobertura": cobertura, "precisao": precisao,
    }


def _xlsx(cliente, pontos):
    r = cliente.post("/api/exportar-tabela-frete-xlsx", json={"hub": {"uf": "DF", "ibge": 5300108}, "pontos_cobertos": pontos})
    assert r.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(r.content))


def test_xlsx_tem_coluna_cobertura_por_ultimo_mantem_parciais_e_aba_de_fonte(cliente):
    wb = _xlsx(cliente, [_ponto("Total"), _ponto("Parcial"), _ponto("Parcial", precisao="faixa")])
    assert wb.sheetnames == ["Prazos e preços", "TZR e TDE", "Cobertura e fonte"]
    ws = wb["Prazos e preços"]
    cabecalho = [c.value for c in ws[3]]
    assert cabecalho[-1] == "Cobertura" and cabecalho[0] == "Código IBGE" and cabecalho[-2] == "ICMS sobre o pedágio"
    linhas = [[c.value for c in row] for row in ws.iter_rows(min_row=4)]
    assert [l[-1] for l in linhas] == ["Total", "Parcial", "Parcial"]  # os parciais continuam na planilha
    assert all(len(l) == len(cabecalho) for l in linhas)  # nenhuma coluna do modelo deslocada
    assert linhas[0][0] == 5300108 and linhas[0][5] == "70254000" and linhas[0][7] == 1


def test_xlsx_legenda_de_uma_linha_e_atribuicao_na_aba_extra(cliente):
    aba = _xlsx(cliente, [_ponto("Total")])["Cobertura e fonte"]
    assert aba["A1"].value == LEGENDA_COBERTURA
    assert aba["A2"].value == "Fonte: IBGE, CNEFE 2022"
    assert aba["A3"].value is None


def test_xlsx_so_com_faixas_manuais_nao_atribui_o_cnefe(cliente):
    aba = _xlsx(cliente, [_ponto("Parcial", precisao="faixa")])["Cobertura e fonte"]
    assert aba["A1"].value == LEGENDA_COBERTURA and aba["A2"].value is None


def test_xlsx_sem_codigo_ibge_deixa_a_celula_vazia_em_vez_de_inventar(cliente):
    r = cliente.post(
        "/api/exportar-tabela-frete-xlsx",
        json={"hub": {"uf": "SC"}, "pontos_cobertos": [{**_ponto("Parcial", precisao="faixa"), "ibge": None}]},
    )
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Prazos e preços"]
    assert ws.cell(row=4, column=1).value is None
