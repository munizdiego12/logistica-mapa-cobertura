"""GZipMiddleware (respostas a partir de ~1 KB), CORS e downloads continuam funcionando, e o log com o tempo da cobertura."""
import asyncio
import gzip
import io
import logging
import random
import re

import openpyxl
import pytest
from fastapi.testclient import TestClient

import main
from backend.tests.fakes_asyncpg import FakePool, criar_banco_sqlite

ORIGEM = "https://meu-frontend.vercel.app"


class OperadorFalso:
    nome = "Maria Souza"
    email = "maria@zubale.com"


def _faixas(n, semente=1):
    random.seed(semente)
    return [{"prefixo": str(10000 + i), "uf": "SP", "cod_municipio": 3550308, "cidade": "São Paulo", "localidade": "CENTRO", "localidade_pct": 100.0,
             "lat": -23.56, "lon": -46.65, "n_enderecos": 120, "dispersao_km": 0.6, "distancia_km": random.random() ** 0.5 * 30} for i in range(n)]


@pytest.fixture
def api(monkeypatch):
    banco = criar_banco_sqlite()
    estado = {"faixas": _faixas(4000), "atraso": 0.0}

    async def get_pool():
        return FakePool(banco)

    async def prefixos(lat, lon, raio):
        if estado["atraso"]:
            await asyncio.sleep(estado["atraso"])
        return estado["faixas"]

    async def geocode_falso(client, *a, **k):
        return -15.8198, -47.9012, "Brasília", "DF", "70347000", "Asa Sul", "VALIDO"

    main.app.dependency_overrides[main.obter_operador_atual] = lambda: OperadorFalso()
    monkeypatch.setattr(main.database, "get_pool", get_pool)
    monkeypatch.setattr(main.database, "consultar_prefixos_por_raio", prefixos)
    monkeypatch.setattr(main, "geocode_async", geocode_falso)
    cliente = TestClient(main.app)
    estado["loja_id"] = cliente.post("/api/lojas", json={"nome": "Atacadão Paulista", "uf": "SP", "endereco": "Av. Paulista, 1000", "lat": "-23,5614", "lng": "-46,6565"}).json()["id"]
    yield cliente, estado
    main.app.dependency_overrides.clear()


# ----------------------------------------------------------------------------------------------------------
# GZip
# ----------------------------------------------------------------------------------------------------------
def test_gzip_esta_ligado_com_minimo_de_cerca_de_1000_bytes():
    gz = [m for m in main.app.user_middleware if m.cls.__name__ == "GZipMiddleware"]
    assert len(gz) == 1 and gz[0].kwargs["minimum_size"] == 1000
    nomes = [m.cls.__name__ for m in main.app.user_middleware]
    assert nomes.index("CORSMiddleware") < nomes.index("GZipMiddleware")  # o CORS continua sendo a camada de fora


def test_cobertura_grande_sai_compactada_e_o_conteudo_e_o_mesmo(api):
    cliente, estado = api
    corpo = {"loja_id": estado["loja_id"], "raio_km": 30}
    com = cliente.post("/api/cobertura-ceps", json=corpo, headers={"Accept-Encoding": "gzip"})
    sem = cliente.post("/api/cobertura-ceps", json=corpo, headers={"Accept-Encoding": "identity"})
    assert com.status_code == sem.status_code == 200
    assert com.headers.get("content-encoding") == "gzip" and "content-encoding" not in sem.headers
    assert "Accept-Encoding" in com.headers.get("vary", "")
    assert com.json() == sem.json() and com.json()["total_pontos"] == 4000
    bruto = gzip.compress(sem.content)
    assert len(sem.content) > 1_000_000 and len(bruto) < len(sem.content) / 5


def test_resposta_pequena_nao_e_compactada(api):
    cliente, estado = api
    r = cliente.get("/api/modais", headers={"Accept-Encoding": "gzip"})
    assert r.status_code == 200 and len(r.content) < 1000
    assert "content-encoding" not in r.headers


def test_download_do_modelo_xlsx_continua_funcionando_com_e_sem_gzip(api):
    cliente, estado = api
    for aceita in ("gzip", "identity"):
        r = cliente.get("/api/modelo-xlsx", headers={"Accept-Encoding": aceita})
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"] and "modelo_pedidos_zubale.xlsx" in r.headers["content-disposition"]
        planilha = openpyxl.load_workbook(io.BytesIO(r.content))
        assert planilha["Modelo_Pedidos"]["A1"].value == "id"


def test_download_da_tabela_xlsx_com_milhares_de_faixas_continua_funcionando_com_gzip(api):
    cliente, estado = api
    cobertura = cliente.post("/api/cobertura-ceps", json={"loja_id": estado["loja_id"], "raio_km": 30}).json()
    corpo = {"hub": cobertura["hub"], "pontos_cobertos": cobertura["pontos_cobertos"]}  # ~1,2 MB enviados de volta, como a tela faz
    resultados = {}
    for aceita in ("gzip", "identity"):
        r = cliente.post("/api/exportar-tabela-frete-xlsx", json=corpo, headers={"Accept-Encoding": aceita})
        assert r.status_code == 200, r.text[:200]
        assert "spreadsheetml" in r.headers["content-type"] and "tabela_frete_completa.xlsx" in r.headers["content-disposition"]
        planilha = openpyxl.load_workbook(io.BytesIO(r.content))
        resultados[aceita] = [aba.max_row for aba in planilha]
    assert resultados["gzip"] == resultados["identity"] and max(resultados["gzip"]) >= 4000  # todas as faixas na planilha


def test_csv_da_tabela_e_gerado_na_propria_tela_entao_nao_passa_pelo_servidor():
    """A exportação CSV é montada no navegador (App.jsx); no servidor só existem os downloads XLSX, cobertos acima."""
    rotas = [r.path for r in main.app.routes if "csv" in getattr(r, "path", "").lower()]
    assert rotas == []


def test_cors_continua_funcionando_na_resposta_comum_na_compactada_e_no_preflight(api):
    cliente, estado = api
    pre = cliente.options("/api/cobertura-ceps", headers={
        "Origin": ORIGEM, "Access-Control-Request-Method": "POST", "Access-Control-Request-Headers": "authorization,content-type"})
    assert pre.status_code == 200
    assert pre.headers["access-control-allow-origin"] in ("*", ORIGEM) and "POST" in pre.headers["access-control-allow-methods"]
    grande = cliente.post("/api/cobertura-ceps", json={"loja_id": estado["loja_id"]}, headers={"Origin": ORIGEM, "Accept-Encoding": "gzip"})
    assert grande.headers.get("content-encoding") == "gzip"
    assert grande.headers["access-control-allow-origin"] in ("*", ORIGEM)
    pequena = cliente.get("/api/modais", headers={"Origin": ORIGEM, "Accept-Encoding": "gzip"})
    assert pequena.headers["access-control-allow-origin"] in ("*", ORIGEM)
    xlsx = cliente.get("/api/modelo-xlsx", headers={"Origin": ORIGEM, "Accept-Encoding": "gzip"})
    assert xlsx.headers["access-control-allow-origin"] in ("*", ORIGEM)


# ----------------------------------------------------------------------------------------------------------
# Log com o tempo da cobertura
# ----------------------------------------------------------------------------------------------------------
def _linhas_de_tempo(caplog):
    return [r.getMessage() for r in caplog.records if r.name == "zubale.cobertura"]


def _numeros(linha):
    return {k: v for k, v in re.findall(r"(\w+)=(\S+)", linha)}


def test_cobertura_por_loja_registra_tempo_total_tempo_do_sql_e_quantidade_de_faixas(api, caplog):
    cliente, estado = api
    estado["atraso"] = 0.05  # a consulta "demora" 50 ms
    caplog.set_level(logging.INFO, logger="zubale.cobertura")
    assert cliente.post("/api/cobertura-ceps", json={"loja_id": estado["loja_id"], "raio_km": 40}).status_code == 200
    linhas = _linhas_de_tempo(caplog)
    assert len(linhas) == 1
    d = _numeros(linhas[0])
    assert d["resultado"] == "ok" and d["faixas"] == "4000" and d["modo"] == "loja" and d["raio_km"] == "40.0"
    assert int(d["sql_ms"]) >= 50 and int(d["total_ms"]) >= int(d["sql_ms"])  # o total inclui a consulta
    assert "total_ms=" in linhas[0] and "sql_ms=" in linhas[0] and "faixas=4000" in linhas[0]


def test_cobertura_por_endereco_tambem_registra(api, caplog):
    cliente, estado = api
    estado["faixas"] = _faixas(120)
    caplog.set_level(logging.INFO, logger="zubale.cobertura")
    r = cliente.post("/api/cobertura-ceps", json={"origem_rua": "SQS 108", "origem_num": "1", "origem_cep": "70347000"})
    assert r.status_code == 200
    d = _numeros(_linhas_de_tempo(caplog)[0])
    assert (d["resultado"], d["faixas"], d["modo"], d["raio_km"]) == ("ok", "120", "endereco", "30.0") and int(d["sql_ms"]) >= 0


def test_o_log_nao_tem_dado_pessoal_nem_string_de_conexao(api, caplog, monkeypatch):
    cliente, estado = api
    monkeypatch.setenv("DATABASE_URL", "postgres://usuario:segredo@host.neon.tech/db?sslmode=require")
    caplog.set_level(logging.INFO, logger="zubale.cobertura")
    cliente.post("/api/cobertura-ceps", json={"loja_id": estado["loja_id"]})
    cliente.post("/api/cobertura-ceps", json={"origem_rua": "SQS 108", "origem_num": "1", "origem_cep": "70347000"})
    texto = "\n".join(_linhas_de_tempo(caplog))
    assert texto
    for proibido in ("Atacadão", "Paulista", "SQS", "70347", "maria", "zubale.com", "Maria Souza", "postgres", "segredo", "neon", "sslmode", "-23,5614", "-23.5614", "-46.6565"):
        assert proibido.lower() not in texto.lower(), proibido
    assert all(re.fullmatch(r"\[cobertura\] resultado=\S+ total_ms=\d+ sql_ms=(\d+|n/d) faixas=\d+ modo=(loja|endereco) raio_km=\S+", l) for l in _linhas_de_tempo(caplog))


def test_falha_do_banco_tambem_deixa_o_tempo_no_log_e_a_resposta_continua_503(api, caplog, monkeypatch):
    cliente, estado = api

    async def quebra(lat, lon, raio):
        await asyncio.sleep(0.02)
        raise RuntimeError("password authentication failed for user neondb_owner")

    monkeypatch.setattr(main.database, "consultar_prefixos_por_raio", quebra)
    caplog.set_level(logging.INFO, logger="zubale.cobertura")
    r = cliente.post("/api/cobertura-ceps", json={"loja_id": estado["loja_id"]})
    assert r.status_code == 503 and "neondb" not in r.text
    linhas = _linhas_de_tempo(caplog)
    d = _numeros(linhas[0])
    assert d["resultado"] == "erro_503" and d["faixas"] == "0" and int(d["sql_ms"]) >= 20
    assert "neondb" not in linhas[0] and "password" not in linhas[0]


def test_pedido_invalido_registra_o_erro_sem_tempo_de_sql(api, caplog):
    cliente, estado = api
    caplog.set_level(logging.INFO, logger="zubale.cobertura")
    assert cliente.post("/api/cobertura-ceps", json={"loja_id": estado["loja_id"], "raio_km": 500}).status_code == 422
    d = _numeros(_linhas_de_tempo(caplog)[0])
    assert d["resultado"] == "erro_422" and d["sql_ms"] == "n/d" and d["faixas"] == "0"
