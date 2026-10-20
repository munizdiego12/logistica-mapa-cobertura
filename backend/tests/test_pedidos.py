"""Etapa 7: pedidos por CSV/XLSX (modelo padrão e modelo antigo), validação linha a linha, casamento da loja, fonte de pedidos com
interface única e privacidade (cliente e endereço nunca em log nem em mensagem de erro)."""
import asyncio
import io
import json
import logging
from datetime import date, time

import openpyxl
import pytest
from fastapi.testclient import TestClient

import casamento_lojas
import main
import pedidos_fonte
from backend.tests.fakes_asyncpg import FakePool, criar_banco_sqlite

CAB = ["id", "loja", "cliente", "endereco", "numero", "complemento", "bairro", "cidade", "uf", "cep", "data_entrega", "janela_inicio", "janela_fim", "volume", "observacao"]


class OperadorFalso:
    nome = "Maria Souza"
    email = "maria@zubale.com"


@pytest.fixture
def banco():
    return criar_banco_sqlite()


@pytest.fixture
def api(monkeypatch, banco):
    async def get_pool():
        return FakePool(banco)

    main.app.dependency_overrides[main.obter_operador_atual] = lambda: OperadorFalso()
    monkeypatch.setattr(main.database, "get_pool", get_pool)
    yield TestClient(main.app), banco
    main.app.dependency_overrides.clear()


def _loja(banco, nome, operacao=None, codigo=None, ativo=1):
    chave = "".join(casamento_lojas._sem_acento(operacao).lower().split()) if operacao else None
    banco.db.execute(
        "INSERT INTO lojas (nome, lat, lng, nome_operacao, nome_operacao_chave, codigo_externo, ativo) VALUES (?, -23.5, -46.6, ?, ?, ?, ?)",
        (nome, operacao, chave, codigo, ativo),
    )


def _csv(linhas, cabecalho=CAB, delimitador=";"):
    saida = [delimitador.join(cabecalho)]
    for linha in linhas:
        saida.append(delimitador.join(str(c) for c in linha))
    return ("\n".join(saida) + "\n").encode("utf-8")


def _pedido(**extra):
    base = dict(id="1", loja="", cliente="Cliente A", endereco="Avenida Paulista", numero="1500", complemento="", bairro="Bela Vista",
                cidade="Sao Paulo", uf="SP", cep="01310-100", data_entrega="2026-10-21", janela_inicio="08:00", janela_fim="13:00", volume="1", observacao="")
    base.update(extra)
    return [base[c] for c in CAB]


def _validar(cliente, conteudo, nome="pedidos.csv"):
    return cliente.post("/api/pedidos/validar", files={"file": (nome, conteudo, "application/octet-stream")})


def _relatorio(cliente, linhas, **kw):
    r = _validar(cliente, _csv(linhas, **kw))
    assert r.status_code == 200, r.text
    return r.json()


def _xlsx(cabecalho, linhas, aba="Pedidos"):
    pasta = openpyxl.Workbook()
    ws = pasta.active
    ws.title = aba
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def _motivos(rel):
    return [m for e in rel["erros"] for m in e["motivos"]]


# ----------------------------------------------------------------------------------------------------------
# Modelos: novo e antigo
# ----------------------------------------------------------------------------------------------------------
def test_modelo_csv_baixado_e_aceito_de_volta(api):
    cliente, _ = api
    r = cliente.get("/api/modelo-csv")
    assert r.status_code == 200 and "modelo_pedidos_zubale.csv" in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf") and r.content.decode("utf-8-sig").splitlines()[0] == ";".join(CAB)
    rel = _validar(cliente, r.content).json()
    assert rel["resumo"]["validas"] == 5 and rel["resumo"]["com_erro"] == 0
    assert rel["pedidos"][4]["volume"] == 1  # volume vazio vale 1
    assert rel["pedidos"][0]["janela_inicio"] == "08:00" and rel["pedidos"][0]["cep"] == "01310-100"
    assert rel["colunas_opcionais_ausentes"] == []


def test_modelo_xlsx_baixado_tem_instrucoes_e_e_aceito_de_volta(api):
    cliente, _ = api
    r = cliente.get("/api/modelo-xlsx")
    pasta = openpyxl.load_workbook(io.BytesIO(r.content))
    assert pasta.sheetnames == ["Modelo_Pedidos", "Instruções"]
    assert [c.value for c in pasta["Modelo_Pedidos"][1]] == CAB
    assert any("Obrigatórias" in str(l[0].value) for l in pasta["Instruções"].iter_rows())
    rel = _validar(cliente, r.content, "modelo.xlsx").json()
    assert rel["resumo"]["validas"] == 5 and rel["resumo"]["com_erro"] == 0 and rel["formato"] == "xlsx"


def test_modelo_antigo_csv_e_xlsx_continuam_aceitos(api):
    cliente, _ = api
    antigo = ["id", "endereco", "numero", "bairro", "cidade", "uf", "cep", "volume"]
    linhas = [[101, "Avenida Paulista", "1500", "Bela Vista", "Sao Paulo", "SP", "01310-100", 1], [102, "Rua Augusta", "1500", "Consolacao", "Sao Paulo", "SP", "01304-001", 2]]
    por_csv = _validar(cliente, _csv(linhas, cabecalho=antigo, delimitador=",")).json()
    por_xlsx = _validar(cliente, _xlsx(antigo, linhas, aba="Modelo_Pedidos"), "antigo.xlsx").json()
    for rel in (por_csv, por_xlsx):
        assert rel["resumo"]["validas"] == 2 and rel["resumo"]["com_erro"] == 0
        assert [p["id"] for p in rel["pedidos"]] == ["101", "102"] and rel["pedidos"][1]["volume"] == 2
        assert set(rel["colunas_opcionais_ausentes"]) >= {"loja", "cliente", "data_entrega", "janela_inicio"}


def test_alias_upload_continua_funcionando_com_total(api):
    cliente, _ = api
    r = cliente.post("/api/upload", files={"file": ("p.csv", _csv([_pedido()]), "text/csv")})
    assert r.status_code == 200 and r.json()["total"] == 1


@pytest.mark.parametrize("cabecalho", [
    ["ID do pedido", "Endereço", "Número", "Cliente", "Cidade", "UF", "CEP", "Data de entrega", "Volume"],
    ["id_do_pedido", "ENDERECO", "NUMERO", "CLIENTE", "CIDADE", "ESTADO", "Cep", "data entrega", "QTD"],
    [" Pedido ", "Logradouro", "Nro", "Nome do cliente", "Município", "uf", "zipcode", "DataDeEntrega", "vol"],
])
def test_cabecalhos_com_variacoes_sao_reconhecidos(api, cabecalho):
    cliente, _ = api
    rel = _relatorio(cliente, [["77", "Avenida Paulista", "10", "Maria", "Sao Paulo", "SP", "01310100", "21/10/2026", "3"]], cabecalho=cabecalho)
    assert rel["resumo"]["validas"] == 1 and rel["resumo"]["com_erro"] == 0, rel["erros"]
    p = rel["pedidos"][0]
    assert (p["id"], p["endereco"], p["numero"], p["cliente"], p["uf"], p["cep"], p["data_entrega"], p["volume"]) == (
        "77", "Avenida Paulista", "10", "Maria", "SP", "01310-100", "2026-10-21", 3)


def test_separadores_e_codificacoes(api):
    cliente, _ = api
    linhas = [_pedido(cliente="José", cidade="São Paulo")]
    assert _validar(cliente, _csv(linhas, delimitador=",")).json()["resumo"]["validas"] == 1
    assert _validar(cliente, _csv(linhas, delimitador="\t")).json()["resumo"]["validas"] == 1
    latin = _csv(linhas).decode("utf-8").encode("cp1252")
    rel = _validar(cliente, latin).json()
    assert rel["pedidos"][0]["cidade"] == "São Paulo" and rel["pedidos"][0]["cliente"] == "José"


def test_celulas_tipadas_do_excel(api):
    cliente, _ = api
    linhas = [[5, "", "Cli", "Rua A", 100, "", "Centro", "Campinas", "SP", 13010000, date(2026, 10, 22), time(8, 0), time(12, 30), 2.0, ""],
              [6, "", "Cli", "Rua B", 10, "", "Centro", "Campinas", "SP", 1310100, date(2026, 10, 22), None, None, None, ""]]
    rel = _validar(cliente, _xlsx(CAB, linhas), "t.xlsx").json()
    assert rel["resumo"]["validas"] == 2, rel["erros"]
    a, b = rel["pedidos"]
    assert (a["id"], a["numero"], a["cep"], a["data_entrega"], a["janela_inicio"], a["janela_fim"], a["volume"]) == ("5", "100", "13010-000", "2026-10-22", "08:00", "12:30", 2)
    assert b["cep"] == "01310-100" and any("7 dígitos" in av for av in b["avisos"])  # o Excel tirou o zero do CEP


def test_xlsx_com_titulo_acima_do_cabecalho(api):
    cliente, _ = api
    pasta = openpyxl.Workbook()
    ws = pasta.active
    ws.append(["Pedidos da semana"])
    ws.append(CAB)
    ws.append(_pedido())
    saida = io.BytesIO()
    pasta.save(saida)
    rel = _validar(cliente, saida.getvalue(), "t.xlsx").json()
    assert rel["resumo"]["validas"] == 1 and rel["pedidos"][0]["linha"] == 3


# ----------------------------------------------------------------------------------------------------------
# Validação linha a linha
# ----------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("mudanca, trecho", [
    ({"id": ""}, "campo obrigatório vazio: id"),
    ({"endereco": ""}, "campo obrigatório vazio: endereco"),
    ({"numero": ""}, "campo obrigatório vazio: numero"),
    ({"cidade": ""}, "campo obrigatório vazio: cidade"),
    ({"uf": ""}, "campo obrigatório vazio: uf"),
    ({"uf": "ZZ"}, "UF inválida"),
    ({"cep": "123"}, "CEP inválido"),
    ({"cep": "abcdefgh"}, "CEP inválido"),
    ({"cep": "0000-0000"}, "CEP inválido"),
    ({"volume": "abc"}, "volume inválido"),
    ({"volume": "0"}, "volume inválido"),
    ({"volume": "-2"}, "volume inválido"),
    ({"volume": "1,5"}, "volume inválido"),
    ({"data_entrega": "31/02/2026"}, "data de entrega inválida"),
    ({"data_entrega": "amanhã"}, "data de entrega inválida"),
    ({"janela_inicio": "25:00"}, "horário de início da janela inválido"),
    ({"janela_fim": "tarde"}, "horário de fim da janela inválido"),
    ({"janela_inicio": "13:00", "janela_fim": "08:00"}, "janela invertida"),
    ({"janela_inicio": "08:00", "janela_fim": "08:00"}, "janela invertida"),
    ({"janela_fim": ""}, "janela incompleta"),
    ({"endereco": "x" * 251}, "texto longo demais em: endereco"),
])
def test_cada_tipo_de_erro_e_reportado_com_linha_e_motivo(api, mudanca, trecho):
    cliente, _ = api
    rel = _relatorio(cliente, [_pedido(id="9"), _pedido(**{"id": "10", **mudanca})])
    assert rel["resumo"] == {"linhas": 2, "validas": 1, "com_erro": 1, "com_aviso": 0, "sem_data": 0, "sem_loja": 0} or rel["resumo"]["com_erro"] == 1
    assert [e["linha"] for e in rel["erros"]] == [3]
    assert any(trecho in m for m in rel["erros"][0]["motivos"]), rel["erros"]
    assert [p["id"] for p in rel["pedidos"]] == ["9"]


def test_varios_erros_na_mesma_linha_aparecem_juntos(api):
    cliente, _ = api
    rel = _relatorio(cliente, [_pedido(uf="ZZ", cep="1", volume="x")])
    assert len(rel["erros"]) == 1 and len(rel["erros"][0]["motivos"]) == 3
    assert rel["erros_por_motivo"] == {"UF inválida": 1, "CEP inválido": 1, "volume inválido": 1}


def test_id_repetido_vale_a_primeira_e_aponta_a_linha(api):
    cliente, _ = api
    rel = _relatorio(cliente, [_pedido(id="A1"), _pedido(id="B2"), _pedido(id="a1"), _pedido(id="A1")])
    assert [p["id"] for p in rel["pedidos"]] == ["A1", "B2"]
    assert [(e["linha"], e["motivos"][0]) for e in rel["erros"]] == [(4, "id repetido (a primeira vez está na linha 2)"), (5, "id repetido (a primeira vez está na linha 2)")]


def test_sem_cep_e_aviso_e_nao_erro(api):
    cliente, _ = api
    rel = _relatorio(cliente, [_pedido(cep="")])
    assert rel["resumo"]["validas"] == 1 and rel["resumo"]["com_aviso"] == 1
    assert rel["avisos"][0]["quantidade"] == 1 and "sem CEP" in rel["avisos"][0]["motivo"] and rel["avisos"][0]["linhas"] == [2]


@pytest.mark.parametrize("texto, esperado", [("2026-10-21", "2026-10-21"), ("21/10/2026", "2026-10-21"), ("21/10/26", "2026-10-21"), ("21-10-2026", "2026-10-21"),
                                             ("2026-10-21 00:00:00", "2026-10-21"), ("1/2/2026", "2026-02-01")])
def test_formatos_de_data(texto, esperado):
    assert pedidos_fonte.interpretar_data(texto) == esperado


@pytest.mark.parametrize("texto, esperado", [("8", "08:00"), ("8:30", "08:30"), ("08:00:00", "08:00"), ("8h", "08:00"), ("8h30", "08:30"), ("23:59", "23:59")])
def test_formatos_de_hora(texto, esperado):
    assert pedidos_fonte.interpretar_hora(texto) == esperado


def test_linhas_vazias_sao_ignoradas_e_numeros_de_linha_sao_do_arquivo(api):
    cliente, _ = api
    conteudo = _csv([_pedido(id="1"), [""] * len(CAB), _pedido(id="2", uf="ZZ")])
    rel = _validar(cliente, conteudo).json()
    assert rel["resumo"]["linhas"] == 2 and rel["erros"][0]["linha"] == 4


def test_datas_e_contagens_do_lote(api):
    cliente, _ = api
    rel = _relatorio(cliente, [_pedido(id="1", data_entrega="2026-10-22"), _pedido(id="2", data_entrega="2026-10-21"), _pedido(id="3", data_entrega="")])
    assert rel["datas"] == ["2026-10-21", "2026-10-22"] and rel["resumo"]["sem_data"] == 1


# ----------------------------------------------------------------------------------------------------------
# Loja: casamento com o cadastro
# ----------------------------------------------------------------------------------------------------------
def test_loja_por_nome_de_operacao_codigo_externo_nome_e_sigla(api):
    cliente, banco = api
    _loja(banco, "Atacadão Mooca", operacao="SDA Mooca", codigo="S-100")
    _loja(banco, "Sam's Club Campinas")
    _loja(banco, "Carrefour Hiper Osasco")
    rel = _relatorio(cliente, [
        _pedido(id="1", loja="sda  MOOCA"), _pedido(id="2", loja="s-100"), _pedido(id="3", loja="Atacadão Mooca"), _pedido(id="4", loja="ATC - Mooca"),
        _pedido(id="5", loja="SAMS Campinas"), _pedido(id="6", loja="CRFO Hiper Osasco"), _pedido(id="7", loja="Atacadao Moóca"),
    ])
    assert rel["resumo"]["com_erro"] == 0, rel["erros"]
    por_id = {p["id"]: p["loja_nome"] for p in rel["pedidos"]}
    assert por_id == {"1": "Atacadão Mooca", "2": "Atacadão Mooca", "3": "Atacadão Mooca", "4": "Atacadão Mooca", "5": "Sam's Club Campinas",
                      "6": "Carrefour Hiper Osasco", "7": "Atacadão Mooca"}
    assert rel["lojas"] == [{"loja_id": 1, "nome": "Atacadão Mooca", "quantidade": 5}, {"loja_id": 3, "nome": "Carrefour Hiper Osasco", "quantidade": 1},
                            {"loja_id": 2, "nome": "Sam's Club Campinas", "quantidade": 1}] or len(rel["lojas"]) == 3


def test_loja_nao_encontrada_ambigua_e_desativada_sao_erros(api):
    cliente, banco = api
    _loja(banco, "Atacadão Norte")
    _loja(banco, "Atacadão - Norte")  # mesmo nome depois da normalização: ambígua
    _loja(banco, "Atacadão Fechada", ativo=0)
    rel = _relatorio(cliente, [_pedido(id="1", loja="Loja Que Não Existe"), _pedido(id="2", loja="Atacadão Norte"), _pedido(id="3", loja="Atacadão Fechada")])
    assert {e["id"]: e["motivos"][0] for e in rel["erros"]} == {
        "1": "loja não encontrada no cadastro", "2": "loja ambígua: o nome casa com mais de uma loja", "3": "loja desativada no cadastro"}
    assert rel["pedidos"] == []


def test_sem_coluna_loja_nao_consulta_o_cadastro_e_o_lote_vale_para_a_loja_escolhida(api, monkeypatch):
    cliente, _ = api

    async def proibido(*a, **k):
        raise AssertionError("não era para consultar o cadastro de lojas")

    monkeypatch.setattr(main, "_com_conexao_lojas", proibido)
    cab = [c for c in CAB if c != "loja"]
    rel = _relatorio(cliente, [[_pedido()[i] for i, c in enumerate(CAB) if c != "loja"]], cabecalho=cab)
    assert rel["resumo"]["validas"] == 1 and rel["pedidos"][0]["loja_id"] is None and "loja" in rel["colunas_opcionais_ausentes"]


def test_loja_em_branco_com_a_coluna_presente_e_aviso(api, banco):
    cliente, banco = api
    _loja(banco, "Atacadão Mooca", operacao="SDA Mooca")
    rel = _relatorio(cliente, [_pedido(id="1", loja="SDA Mooca"), _pedido(id="2", loja="")])
    assert rel["resumo"]["validas"] == 2 and rel["resumo"]["sem_loja"] == 1
    assert any("vale para a loja escolhida" in a["motivo"] for a in rel["avisos"])
    assert [p["loja_id"] for p in rel["pedidos"]] == [1, None]


def test_casar_loja_e_a_mesma_regra_da_carga_de_motoristas():
    cad = [{"id": 1, "nome": "Atacadão Porto Alegre Humaitá", "nome_operacao": "POA Humaita", "codigo_externo": "X9", "ativo": True},
           {"id": 2, "nome": "Sam's Club Campinas", "ativo": True}]
    assert casamento_lojas.casar_loja("poa  humaitá", cad)[1]["id"] == 1
    assert casamento_lojas.casar_loja("x9", cad)[1]["id"] == 1
    assert casamento_lojas.casar_loja("ATC Porto Alegre Humaita", cad)[1]["id"] == 1
    assert casamento_lojas.casar_loja("SAMs Campinas", cad)[1]["id"] == 2
    assert casamento_lojas.casar_loja("Nada", cad) == ("nao_encontrada", None)
    from backend import carga_motoristas
    assert carga_motoristas.chave_da_loja("ATC Av das Torres") == casamento_lojas.chave_da_loja("ATC Av das Torres") == "atacadao avenida das torres"


# ----------------------------------------------------------------------------------------------------------
# Arquivo com problema
# ----------------------------------------------------------------------------------------------------------
def test_faltando_coluna_obrigatoria_recusa_o_arquivo_com_mensagem_clara(api):
    cliente, _ = api
    r = _validar(cliente, _csv([["1", "Rua A"]], cabecalho=["id", "endereco"]))
    assert r.status_code == 400
    assert r.json()["detail"].startswith("Faltam colunas obrigatórias no cabeçalho: numero, cidade, uf.") and "modelo padrão" in r.json()["detail"]


@pytest.mark.parametrize("nome, conteudo, trecho", [
    ("p.txt", b"id;endereco", "Formato de arquivo inválido"),
    ("p.xls", b"qualquer", ".xls antigas"),
    ("p.csv", b"", "arquivo está vazio"),
    ("p.xlsx", b"isto nao e um xlsx", "Não consegui ler a planilha"),
])
def test_arquivos_invalidos_dao_mensagem_clara(api, nome, conteudo, trecho):
    cliente, _ = api
    r = _validar(cliente, conteudo, nome)
    assert r.status_code == 400 and trecho in r.json()["detail"]


def test_arquivo_grande_demais_e_recusado(api, monkeypatch):
    cliente, _ = api
    monkeypatch.setattr(pedidos_fonte, "MAX_BYTES", 200)
    r = _validar(cliente, _csv([_pedido(id=str(i)) for i in range(20)]))
    assert r.status_code == 400 and "grande demais" in r.json()["detail"]


def test_celula_gigante_no_csv_da_mensagem_fixa(api):
    cliente, _ = api
    r = _validar(cliente, ("id;endereco;numero;cidade;uf\n1;" + "x" * 200000 + ";1;SP;SP\n").encode())
    assert r.status_code == 400 and "Não consegui ler o CSV" in r.json()["detail"]


def test_rotas_exigem_login():
    cliente = TestClient(main.app)
    assert _validar(cliente, _csv([_pedido()])).status_code == 401
    assert cliente.get("/api/modelo-csv").status_code == 401
    assert cliente.get("/api/modelo-xlsx").status_code == 401


# ----------------------------------------------------------------------------------------------------------
# Fonte de pedidos com interface única
# ----------------------------------------------------------------------------------------------------------
class FonteDeTeste:
    """Uma fonte futura (Sheets, Superset, BigQuery) só precisa devolver um LoteBruto com os nomes de campo padrão."""

    def ler(self):
        campos = dict(zip(CAB, _pedido(id="S1")))
        return pedidos_fonte.LoteBruto(origem="teste", formato="memoria", colunas={c: c for c in CAB},
                                       linhas=[pedidos_fonte.LinhaBruta(2, campos), pedidos_fonte.LinhaBruta(3, {**campos, "id": "S1"})])


def test_carregar_lote_usa_a_mesma_validacao_para_qualquer_fonte():
    rel = pedidos_fonte.carregar_lote(FonteDeTeste())
    assert rel.origem == "teste" and len(rel.pedidos) == 1 and rel.erros[0]["motivos"][0].startswith("id repetido")
    assert "arquivo" in pedidos_fonte.FONTES and pedidos_fonte.FONTES["arquivo"] is pedidos_fonte.FonteArquivo
    rel_arquivo = pedidos_fonte.carregar_lote(pedidos_fonte.FonteArquivo(_csv([_pedido()]), "x.csv"))
    assert rel_arquivo.origem == "arquivo" and len(rel_arquivo.pedidos) == 1


# ----------------------------------------------------------------------------------------------------------
# Privacidade: cliente e endereço nunca em log nem em mensagem de erro
# ----------------------------------------------------------------------------------------------------------
SEGREDOS = ("Fulano Segredoso", "Rua Sigilosa", "Complemento Reservado", "Bairro Escondido")


def _arquivo_com_dados_pessoais():
    return _csv([
        _pedido(id="1", cliente=SEGREDOS[0], endereco=SEGREDOS[1] + " 999", complemento=SEGREDOS[2], bairro=SEGREDOS[3], uf="ZZ"),
        _pedido(id="2", cliente=SEGREDOS[0], endereco=SEGREDOS[1] + " 1000", cep="12", volume="x", janela_inicio="9", janela_fim="8"),
        _pedido(id="3", cliente=SEGREDOS[0], endereco=SEGREDOS[1] + " 1", loja="Loja Fantasma"),
        _pedido(id="3", cliente=SEGREDOS[0], endereco=SEGREDOS[1] + " 2", cep=""),
    ])


def test_relatorio_de_erros_nao_traz_cliente_nem_endereco(api, caplog, capsys):
    cliente, _ = api
    caplog.set_level(logging.DEBUG)
    rel = _validar(cliente, _arquivo_com_dados_pessoais()).json()
    assert rel["resumo"]["com_erro"] >= 3
    parte_de_erros = json.dumps({"erros": rel["erros"], "motivos": rel["erros_por_motivo"], "avisos": rel["avisos"]}, ensure_ascii=False)
    for segredo in SEGREDOS:
        assert segredo not in parte_de_erros
        assert segredo not in caplog.text
        assert segredo not in capsys.readouterr().out


def test_erro_de_leitura_nao_vaza_o_conteudo_do_arquivo(api, monkeypatch, caplog, capsys):
    cliente, _ = api
    caplog.set_level(logging.DEBUG)

    def estoura(self):
        raise RuntimeError(f"falha em {SEGREDOS[1]} 999 de {SEGREDOS[0]}")

    monkeypatch.setattr(pedidos_fonte.FonteArquivo, "ler", estoura)
    r = _validar(cliente, _arquivo_com_dados_pessoais())
    assert r.status_code == 400 and r.json()["detail"] == main.MSG_ARQUIVO_ILEGIVEL
    saida = capsys.readouterr().out
    assert "RuntimeError" in saida
    for segredo in SEGREDOS:
        assert segredo not in r.text and segredo not in saida and segredo not in caplog.text


def test_erro_da_roteirizacao_nao_traz_o_endereco_do_pedido(monkeypatch, capsys, caplog):
    async def geocodificar_quebrado(pedidos):
        raise RuntimeError(f"não achei {SEGREDOS[1]} 999")

    monkeypatch.setattr(main, "geocodificar_lote", geocodificar_quebrado)
    requisicao = main.OtimizarRequest(pedidos=[{"id": "1", "Endereco": SEGREDOS[1], "Numero": "999", "Cidade": "Sao Paulo", "UF": "SP"}])
    job_id = main._novo_job()
    loja = {"id": 1, "nome": "Loja", "endereco": None, "lat": -23.5, "lng": -46.6}
    asyncio.run(main._processar_otimizacao(job_id, requisicao, [], None, loja))
    job = main.JOBS[job_id]
    assert job["status"] == "erro" and "RuntimeError" in job["erro"]
    assert SEGREDOS[1] not in job["erro"] and SEGREDOS[1] not in capsys.readouterr().out and SEGREDOS[1] not in caplog.text


def test_httpx_nao_registra_a_url_com_o_endereco():
    assert logging.getLogger("httpx").getEffectiveLevel() >= logging.WARNING
    assert logging.getLogger("httpcore").getEffectiveLevel() >= logging.WARNING


def test_o_modulo_nao_guarda_nada_nem_fala_com_banco_de_pedidos():
    fonte = open(pedidos_fonte.__file__, encoding="utf-8").read()
    assert "INSERT" not in fonte.upper() and "asyncpg" not in fonte and "sqlalchemy" not in fonte.lower() and "print(" not in fonte
